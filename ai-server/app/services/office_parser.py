"""DOCX / PPTX 텍스트 추출."""
import io
import re
import zipfile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# 압축 해제 후 총 크기 상한 (zip bomb 방어)
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024


class OfficeParseError(ValueError):
    """Office 문서를 읽을 수 없음 (손상, 암호 문서, 과대 크기 등)."""


def extract_docx(data: bytes) -> str:
    """본문 순서대로 문단과 표를 추출한다. 제목 스타일은 마크다운 '#'로 표시한다."""
    _check_zip(data)
    try:
        doc = Document(io.BytesIO(data))
        blocks = []
        for item in doc.iter_inner_content():
            if isinstance(item, Paragraph):
                text = item.text.strip()
                if text:
                    blocks.append(_heading_prefix(item) + text)
            elif isinstance(item, Table):
                table = _table_text(item.rows, lambda cell: cell.text)
                if table:
                    blocks.append(table)
    except OfficeParseError:
        raise
    except Exception as exc:  # 손상된 XML 등 라이브러리 예외가 다양하다
        raise OfficeParseError(f"DOCX를 읽을 수 없습니다: {exc}") from exc
    return "\n\n".join(blocks)


def extract_pptx(data: bytes) -> str:
    """슬라이드 순서대로 텍스트, 표, 발표자 노트를 추출한다."""
    _check_zip(data)
    try:
        slides = []
        for number, slide in enumerate(Presentation(io.BytesIO(data)).slides, start=1):
            parts = []
            for shape in _iter_shapes(slide.shapes):
                if shape.has_table:
                    table = _table_text(shape.table.rows, lambda cell: cell.text)
                    if table:
                        parts.append(table)
                elif shape.has_text_frame and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text.strip())
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"[발표자 노트]\n{notes}")
            if parts:
                slides.append(f"[슬라이드 {number}]\n" + "\n".join(parts))
    except OfficeParseError:
        raise
    except Exception as exc:
        raise OfficeParseError(f"PPTX를 읽을 수 없습니다: {exc}") from exc
    return "\n\n".join(slides)


def _check_zip(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise OfficeParseError(
            "손상되었거나 올바른 Office 문서가 아닙니다. (암호가 걸린 문서는 지원하지 않습니다)"
        ) from exc
    if total > MAX_UNCOMPRESSED_BYTES:
        raise OfficeParseError("문서의 압축 해제 크기가 너무 큽니다.")


def _heading_prefix(paragraph: Paragraph) -> str:
    style = paragraph.style.name if paragraph.style is not None else ""
    if style == "Title":
        return "# "
    match = re.fullmatch(r"Heading (\d)", style or "")
    return "#" * min(int(match.group(1)) + 1, 6) + " " if match else ""


def _table_text(rows, cell_text) -> str:
    lines = []
    for row in rows:
        cells, seen = [], set()
        for cell in row.cells:
            # 병합된 셀은 같은 셀 객체가 반복되므로 한 번만 쓴다
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells.append(" ".join(cell_text(cell).split()))
        if any(cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _iter_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape
